# -*- coding: utf-8 -*-
"""核心：客户表格解析、规格归一化、匹配引擎。"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

import xlrd
from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# 文本归一化：把“看起来差不多”的写法统一成同一种
# ---------------------------------------------------------------------------
PHI_CHARS = str.maketrans({"φ": "Φ", "∅": "Φ", "Ø": "Φ", "ø": "Φ", "⌀": "Φ"})
SPEC_HEADER_KEYS = ("采购规格", "规格型号", "规格", "型号")


def norm_text(s) -> str:
    """统一全半角、Φ 各种写法、×/＊ 乘号写法，去掉所有空格。"""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s)).strip()
    s = s.translate(PHI_CHARS)
    s = s.replace("×", "x").replace("*", "x").replace("＊", "x")
    s = re.sub(r"\s+", "", s)
    return s


def search_norm(s) -> str:
    """搜索用的再归一化：在 `norm_text` 之上把标点/符号全去掉，只留汉字与字母数字。

    ★ 为什么要这一层：V3 的拼音码（`Hpdan.Py`）长这样 —— `LS（BT）`、`BXGNLJ2.5*20`。
      `norm_text` 只统一全半角与乘号，括号还在 ⇒ 人打「lsbt」搜"螺丝（杯头）"就匹配不上，
      而他恰恰是按拼音码打的。编号里的点、规格里的斜杠同理。
    ★ 只用于搜索串与关键词，**不参与相似度打分**（打分要保留规格差异，去掉符号会把
      M6/20 与 M6-20 判成同一个东西，那是另一类错）。
    ★ 一律转小写：V3 存的拼音码是大写（`LS（BT）`），而人在手机键盘上打的是小写 `lsbt`
      —— 不统一大小写就等于"拼音码搜索只对一半"。
    """
    return re.sub(r"[^0-9a-z一-鿿]", "", norm_text(s).lower())


@dataclass
class PartRow:
    """客户表格里的一行零件。"""
    source_file: str
    sheet: str
    mold_no: str
    seq: str
    part_no_orig: str      # 客户给的零件编号（基本是空的，保留原样）
    name: str
    material: str
    qty: float
    spec: str
    actual: str            # 实际尺寸
    vendor: str
    heat: str
    remark: str

    erp_code: str = ""     # 匹配到的 hpbh（货品编号）
    erp_unit: str = ""     # 匹配到的 hpdanunit（单位）
    sys_name: str = ""     # 系统产品名（匹配记录的 HpmcKh）
    sys_spec: str = ""     # 系统规格（匹配记录的 GgxhKh）
    status: str = "未匹配"  # 未匹配/记忆匹配/自动匹配/手动匹配
    source: str = ""

    @property
    def merge(self) -> str:
        """客户产品名称（主）：零件名称+采购规格 合并。"""
        return (self.name or "") + (self.spec or "")

    @property
    def merge_a_n(self) -> str:
        """匹配键 A：归一化后的 零件名称+采购规格。"""
        return norm_text(self.merge)

    @property
    def merge_b(self) -> str:
        """匹配键 B（备用）：零件名称+实际尺寸 合并。"""
        return (self.name or "") + (self.actual or "")

    @property
    def merge_b_n(self) -> str:
        """匹配键 B：归一化后的 零件名称+实际尺寸。"""
        return norm_text(self.merge_b)

    @property
    def merge_n(self) -> str:
        """兼容旧逻辑：默认取 A 键。"""
        return self.merge_a_n


# ---------------------------------------------------------------------------
# 读取 Excel（同时支持老版 .xls 和新版 .xlsx）
# ---------------------------------------------------------------------------
def _to_qty(v) -> float:
    if v is None or v == "":
        return 0
    try:
        f = float(v)
        return int(f) if f == int(f) else f
    except (TypeError, ValueError):
        return 0


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _read_matrix(path: Path) -> list[tuple[str, list[list]]]:
    """返回 [(sheet名, [[单元格值...], ...]), ...]"""
    if path.suffix.lower() == ".xls":
        wb = xlrd.open_workbook(str(path))
        out = []
        for sh in wb.sheets():
            matrix = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    v = sh.cell_value(r, c)
                    if sh.cell_type(r, c) == xlrd.XL_CELL_DATE:
                        try:
                            v = xlrd.xldate_as_datetime(v, wb.datemode)
                        except Exception:
                            pass
                    row.append(v)
                matrix.append(row)
            out.append((sh.name, matrix))
        return out
    wb = load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        matrix = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        out.append((ws.title, matrix))
    wb.close()
    return out


def _find_col(header_row: list, keyword: str) -> int | None:
    for i, cell in enumerate(header_row):
        if keyword in _cell_to_str(cell):
            return i
    return None


def _extract_mold_no(matrix: list[list], header_idx: int) -> str:
    for r in range(min(header_idx, 6)):
        row = matrix[r]
        for j, cell in enumerate(row):
            s = _cell_to_str(cell)
            if s.startswith("模号"):
                m = re.split(r"[：:]", s, maxsplit=1)
                if len(m) == 2 and m[1].strip():
                    return m[1].strip()
                for k in range(j + 1, len(row)):
                    v = _cell_to_str(row[k])
                    if v and v not in ("设计", "日期", "设计：", "日期："):
                        return v
    return ""


def _parse_sheet(matrix: list[list], fname: str, sheet_name: str) -> list[PartRow]:
    maxc = max((len(r) for r in matrix), default=0)
    matrix = [list(r) + [""] * (maxc - len(r)) for r in matrix]

    header_idx = None
    for i, row in enumerate(matrix):
        cells = [_cell_to_str(c) for c in row]
        has_spec = any(any(k in c for k in SPEC_HEADER_KEYS) for c in cells)
        if any("零件名称" in c for c in cells) and has_spec:
            header_idx = i
            break
    if header_idx is None:
        return []

    header = matrix[header_idx]
    col = {}
    for key, kw in (
        ("name", "零件名称"),
        ("actual", "实际尺寸"),
        ("qty", "数量"),
        ("material", "材料"),
        ("heat", "热处理"),
        ("vendor", "厂家"),
        ("remark", "备注"),
        ("seq", "序号"),
        ("part_no", "零件编号"),
    ):
        col[key] = _find_col(header, kw)
    col["spec"] = None
    for kw in SPEC_HEADER_KEYS:
        idx = _find_col(header, kw)
        if idx is not None:
            col["spec"] = idx
            break
    if col["name"] is None or col["spec"] is None:
        return []

    mold_no = _extract_mold_no(matrix, header_idx)
    rows: list[PartRow] = []
    for r in range(header_idx + 1, len(matrix)):
        row = matrix[r]
        name = _cell_to_str(row[col["name"]]) if col["name"] is not None else ""
        if not name:
            continue
        rows.append(
            PartRow(
                source_file=fname,
                sheet=sheet_name,
                mold_no=mold_no,
                seq=_cell_to_str(row[col["seq"]]) if col["seq"] is not None else "",
                part_no_orig=_cell_to_str(row[col["part_no"]]) if col["part_no"] is not None else "",
                name=name,
                material=_cell_to_str(row[col["material"]]) if col["material"] is not None else "",
                qty=_to_qty(row[col["qty"]] if col["qty"] is not None else None),
                spec=_cell_to_str(row[col["spec"]]) if col["spec"] is not None else "",
                actual=_cell_to_str(row[col["actual"]]) if col["actual"] is not None else "",
                vendor=_cell_to_str(row[col["vendor"]]) if col["vendor"] is not None else "",
                heat=_cell_to_str(row[col["heat"]]) if col["heat"] is not None else "",
                remark=_cell_to_str(row[col["remark"]]) if col["remark"] is not None else "",
            )
        )
    return rows


def parse_file(path: str | Path) -> list[PartRow]:
    """解析一个客户零件清单文件，返回所有零件行。"""
    path = Path(path)
    rows: list[PartRow] = []
    for sheet_name, matrix in _read_matrix(path):
        rows.extend(_parse_sheet(matrix, path.name, sheet_name))
    return rows


IMPORT_NAME_KEYS = ("货品名称", "物料名称")
IMPORT_CUST_KEYS = ("客户品名", "客户产品名称", "客户名称")
IMPORT_ACTUAL_KEYS = ("实际尺寸", "尺寸")


def parse_import_sheet(matrix: list[list], sheet_name: str = "") -> list[dict]:
    """解析**一张表**的"货品名称|客户品名|采购规格|实际尺寸"。

    ★ 与 `_parse_sheet`（零件清单）分开：这是「导入燕秀品名」用的模板格式，
      表头要求是「货品名称（或物料名称）」+「客户品名（或客户产品名称/客户名称）」。
    """
    header_idx = None
    for i, row in enumerate(matrix):
        cells = [_cell_to_str(c) for c in row]
        has_goods = any(any(k in c for k in IMPORT_NAME_KEYS) for c in cells)
        has_cust = any(any(k in c for k in IMPORT_CUST_KEYS) for c in cells)
        if has_goods and has_cust:
            header_idx = i
            break
    if header_idx is None:
        return []

    header = matrix[header_idx]

    def find(keys: tuple[str, ...]) -> int | None:
        for kw in keys:
            idx = _find_col(header, kw)
            if idx is not None:
                return idx
        return None

    ci = {
        "goods": find(IMPORT_NAME_KEYS),
        "cust": find(IMPORT_CUST_KEYS),
        "spec": find(SPEC_HEADER_KEYS),
        "actual": find(IMPORT_ACTUAL_KEYS),
    }
    if ci["goods"] is None or ci["cust"] is None:
        return []
    out: list[dict] = []
    for r in range(header_idx + 1, len(matrix)):
        row = matrix[r]
        goods = _cell_to_str(row[ci["goods"]]) if ci["goods"] < len(row) else ""
        cust = _cell_to_str(row[ci["cust"]]) if ci["cust"] < len(row) else ""
        if not goods and not cust:
            continue
        out.append({
            "goods_name": goods,
            "cust_name": cust,
            "spec": _cell_to_str(row[ci["spec"]]) if ci["spec"] is not None and ci["spec"] < len(row) else "",
            "actual": _cell_to_str(row[ci["actual"]]) if ci["actual"] is not None and ci["actual"] < len(row) else "",
            "sheet": sheet_name,
        })
    return out


def parse_import_sheets(sheets: list[dict]) -> list[dict]:
    """给 op 用：`[{name, rows}]`（App 传上来的矩阵）→ 导入行。

    ★ App 侧仍然只把文件读成矩阵，**表头识别与取值都在这里**（A1.5 同一条原则）。
    """
    out: list[dict] = []
    for s in sheets or []:
        if not isinstance(s, dict):
            continue
        matrix = [[_cell_to_str(c) for c in (row or [])] for row in (s.get("rows") or [])]
        out.extend(parse_import_sheet(matrix, str(s.get("name") or "")))
    return out


def parse_import_rows(path: str | Path) -> list[dict]:
    """解析"货品名称|客户品名|采购规格|实际尺寸"格式的导入表（桌面端按路径读）。"""
    path = Path(path)
    out: list[dict] = []
    for sheet_name, matrix in _read_matrix(path):
        out.extend(parse_import_sheet(matrix, sheet_name))
    return out


# ---------------------------------------------------------------------------
# 匹配引擎
# ---------------------------------------------------------------------------
def _tokens(s: str) -> set[str]:
    return set(re.split(r"[^0-9A-Za-z\u4e00-\u9fff]+", s)) - {""}


def _token_overlap(a: str, b: str) -> float:
    ta, tb = _tokens(a), _tokens(b)
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / max(len(ta), len(tb))


class Matcher:
    """物料库 + 匹配逻辑。

    items 里的每条物料：
    - code: hpbh（货品编号）
    - unit: HpUnitName（单位）
    - name: hpmckh（货品名称-客户）
    - spec_a: ggxhkh（规格-客户），spec_b: hpbhkh（客户编号/尺寸）
    - key_a: 归一化 hpmckh + ggxhkh
    - key_b: 归一化 hpmckh + hpbhkh
    """

    def __init__(self, storage):
        self.storage = storage
        self.items: list[dict] = []
        self.sys_map: dict[str, dict] = {}   # hpbh -> {sys_name, sys_spec}
        self.source_label = "未加载"

    def load_from_storage(self):
        """从本地 Excel 导入的物料档案加载（备用数据源）。"""
        self.items = self.storage.all_erp_items()
        self.source_label = f"Excel物料档案 {len(self.items)} 条"

    def load_hpdan(self, rows: list[dict], sys_map: dict[str, dict] | None = None):
        """加载 SQL Server hpdan 抓取结果。"""
        self.items = rows
        self.sys_map = sys_map or {}
        self.source_label = f"SQL Server HPKhType {len(rows)} 条"

    def forget(self, reason: str = "未加载"):
        """把物料库整份丢掉（★ 换账套必须调它）。

        ★ 为什么不能留着：`items` 是从**某一本书**里抓出来的 HPKhType。
          切到账套 B 之后还拿着 A 的库给候选，人看到的是"A 账的编号 + B 账的单子"，
          点下去写进去的是一个在 B 账里不存在、或根本是另一种东西的编号 ——
          界面上一切都"正常"，只有对账时才看得出来。
        """
        self.items = []
        self.sys_map = {}
        self.source_label = reason

    def fill_sys(self, row: PartRow):
        """按 hpbh 从物料库/hpdan 补系统产品名、规格。"""
        item = next((it for it in self.items if it.get("code") == row.erp_code), None)
        if item:
            row.sys_name = item.get("sys_name") or item.get("name") or ""
            row.sys_spec = item.get("sys_spec") or item.get("spec_a") or ""
        info = self.sys_map.get(row.erp_code)
        if info:
            if info.get("sys_name"):
                row.sys_name = info["sys_name"]
            if info.get("sys_spec"):
                row.sys_spec = info["sys_spec"]

    def _apply(self, row: PartRow, item: dict, status: str, source: str = ""):
        row.erp_code = item.get("code") or ""
        row.erp_unit = item.get("unit") or ""
        self.fill_sys(row)
        row.status = status
        row.source = source or item.get("name") or ""

    def auto_match(self, row: PartRow) -> None:
        a_n = row.merge_a_n
        b_n = row.merge_b_n
        actual_n = norm_text(row.actual)
        if not (a_n or b_n):
            row.status = "未匹配"
            row.source = ""
            return

        # 第 1 层：历史记忆（同一客户写法再次出现时，直接命中）
        # 键 B 只有在客户行确有“实际尺寸”时才参与，避免只剩名称造成串规格
        for key, label, need_actual in ((a_n, "A", False), (b_n, "B", True)):
            if not key:
                continue
            if need_actual and not actual_n:
                continue
            mem = self.storage.find_memory(key)
            if mem:
                row.erp_code = mem["code"] or ""
                row.erp_unit = mem["unit"] or ""
                self.fill_sys(row)
                row.status = "记忆匹配"
                row.source = f"历史确认-{label}(第{mem['match_count'] + 1}次)"
                self.storage.bump_memory(mem["id"])
                return

        # 第 2 层：方法A 零件名称+采购规格 = hpmckh + ggxhkh
        if a_n:
            exact = [it for it in self.items if it.get("key_a") == a_n]
            if exact:
                best = next((it for it in exact if it.get("code")), exact[0])
                self._apply(row, best, "自动匹配", "规格匹配")
                return

        # 第 3 层：方法B 零件名称+实际尺寸 = hpmckh + hpbhkh
        # 安全条件：客户行“实际尺寸”和物料“HpbhKh”都非空才匹配；
        # 任一侧为空时键B会退化成只有名称，把同名称不同规格的行全部配错。
        if b_n and actual_n:
            exact = [
                it for it in self.items
                if it.get("key_b") == b_n and (it.get("spec_b") or "").strip()
            ]
            if exact:
                best = next((it for it in exact if it.get("code")), exact[0])
                self._apply(row, best, "自动匹配", "实际尺寸匹配")
                return

        # 不做自动近似匹配：避免 M6x16 与 M6x160 这种配错，交给人工确认
        row.status = "未匹配"
        row.source = ""

    def _sys_name(self, it: dict) -> str:
        """物料对应的系统品名（hpdan 里查到的）。搜索时要连它一起命中。"""
        info = (getattr(self, "sys_map", None) or {}).get(it.get("code") or "") or {}
        return info.get("sys_name") or it.get("sys_name") or ""

    def _blob(self, it: dict) -> str:
        """一行物料"能被搜到的字"：客户品名 + 两个规格 + 编号 + 系统品名 + **拼音码**。

        ★ 系统品名一起进来（2026-10-13）：桌面弹窗的筛选框本来就带它，
          而 op 侧原来不带 ⇒ 同一段字，手机上搜不到、电脑上搜得到。
        ★ 拼音码（`Hpdan.Py`，2026-10-13 加）：V3 自己就靠这一列搜货品，
          老板在桌面端打的是 `lsbt`、`bxgnlj`，到手机上打同样的字必须搜得到。
          库里 4.8 万条有 99.99% 填了它，不是冷门字段。
        ★ 整串走 `search_norm`：拼音码长这样 `LS（BT）`，不剥括号打 `lsbt` 就对不上。
        """
        return search_norm(
            (it.get("name_n") or "") + (it.get("spec_a_n") or "")
            + (it.get("spec_b_n") or "") + (it.get("code") or "")
            + self._sys_name(it) + self._sys_py(it))

    def _sys_py(self, it: dict) -> str:
        """这条物料对应系统档案的拼音码（V3 `Hpdan.Py`）。没有就是空串。"""
        info = (getattr(self, "sys_map", None) or {}).get(it.get("code") or "") or {}
        return info.get("sys_py") or it.get("sys_py") or ""

    def _sim(self, row: PartRow, it: dict) -> float:
        """这一行与这条物料的相似度（0~1）。"""
        a_n = row.merge_a_n
        b_n = row.merge_b_n
        name_n = norm_text(row.name)
        spec_n = norm_text(row.spec)
        actual_n = norm_text(row.actual)
        # ★ 2026-10-09 修：不再让"空对空"参与打分。
        #   客户行没有实际尺寸时 b_n 只剩名称；物料 HpbhKh 为空时 key_b 也只剩名称
        #   ⇒ (b_n, key_b) 这一对必然满分，把同名不同规格的候选全顶到 100%
        #   （现场看到"杯头螺丝 M6*35 的两条候选都是 100%"，就是它）。
        #   同理，两侧都没有规格信息时 (a_n, key_a) 也只是"名称比名称"——
        #   那件事已经由下面的 nsim 负责，不该再算一次。
        it_spec_a = it.get("spec_a_n") or ""
        it_spec_b = it.get("spec_b_n") or ""
        pairs = []
        if a_n and it.get("key_a") and (spec_n or it_spec_a):
            pairs.append((a_n, it.get("key_a")))
        if b_n and it.get("key_b") and actual_n and it_spec_b:
            pairs.append((b_n, it.get("key_b")))
        if b_n and it.get("key_a") and actual_n and it_spec_a:
            pairs.append((b_n, it.get("key_a")))
        if a_n and it.get("key_b") and spec_n and it_spec_b:
            pairs.append((a_n, it.get("key_b")))
        msim = 0.0
        tok = 0.0
        for x, y in pairs:
            if x and y:
                msim = max(msim, SequenceMatcher(None, x, y).ratio())
                tok = max(tok, _token_overlap(x, y))
        nsim = SequenceMatcher(None, name_n, it.get("name_n") or "").ratio()
        return msim * 0.6 + tok * 0.2 + nsim * 0.2

    def suggest(self, row: PartRow, keyword: str = "", top: int | None = None,
                search: bool = False) -> list[tuple[float, dict]]:
        """模糊搜索物料库，返回 (相似度, 物料) 候选。

        默认（`search=False`）返回全部达标（相似度≥0.35）的候选并按分数从高到低排序；
        传入 top 时只取前 top 条。

        ★ `search=True` 是**在整库里搜**（App 候选框顶部那栏搜索用）：
          只要关键词命中就返回，**不套 0.35 那道门槛**。
          原来 keyword 是在"已经达标的候选"里再筛一遍 ⇒ 拿编号或系统品名去搜一个
          规格完全不同的东西，会因为"不相似"被丢掉，搜索框形同虚设。
          相似度仍然带回去，只用来排序和显示，不再当筛子。
        """
        # ★ 关键词与搜索串必须过**同一个**归一化：只剥一边的话括号/点会永远对不上
        kw = search_norm(keyword)
        scored: list[tuple[float, dict]] = []
        for it in self.items:
            if kw and kw not in self._blob(it):
                continue
            score = self._sim(row, it)
            if not search and score < 0.35:
                continue
            scored.append((score, it))
        if search:
            # 同分时按编号排，保证同一个关键词每次刷出来的顺序一样（列表会跳很碍点选）
            scored.sort(key=lambda t: (-t[0], str(t[1].get("code") or "")))
            return scored[: (top or 50)]
        scored.sort(key=lambda t: t[0], reverse=True)
        if top:
            return scored[:top]
        return scored
